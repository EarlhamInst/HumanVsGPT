"""Loading manifest workbooks into a typed, comparison-ready model.

Reading these workbooks is not quite trivial. Headers carry an '(optional)'
marker that is presentation, not identity. Sheets are padded with hundreds of
blank rows (one human manifest declares 1000 rows in `person` and populates 19).
Numeric cells arrive from openpyxl as floats, so a NovaSeq model number reads
'6000.0'. Each of those would corrupt a naive comparison, so all three are
resolved here, once, at the boundary.

The loader also records provenance for every file it reads: a content hash and
the workbook's own authoring metadata. A reproducibility claim needs to name the
exact bytes a result came from.
"""

from __future__ import annotations

import datetime as _dt
import hashlib
import re
import zipfile
from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import load_workbook

from .normalise import normalise_column
from .schema import SCORED_SHEETS

_CORE_XML = "docProps/core.xml"
_CORE_FIELDS = {
    "created": "dcterms:created",
    "modified": "dcterms:modified",
    "creator": "dc:creator",
    "last_modified_by": "cp:lastModifiedBy",
}


@dataclass(frozen=True)
class Provenance:
    """Identity of a source file, sufficient to reproduce a result from it."""

    path: str
    sha256: str
    size_bytes: int
    mtime: str
    created: str = ""
    modified: str = ""
    creator: str = ""
    last_modified_by: str = ""

    def as_dict(self) -> dict[str, str | int]:
        return {
            "path": self.path,
            "sha256": self.sha256,
            "size_bytes": self.size_bytes,
            "mtime": self.mtime,
            "xlsx_created": self.created,
            "xlsx_modified": self.modified,
            "xlsx_creator": self.creator,
            "xlsx_last_modified_by": self.last_modified_by,
        }


@dataclass
class Sheet:
    """One worksheet: ordered columns and populated rows keyed by column name."""

    name: str
    columns: list[str]
    rows: list[dict[str, str]] = field(default_factory=list)

    def column(self, name: str) -> list[str]:
        """All values in a column, in row order; '' where absent."""
        return [row.get(name, "") for row in self.rows]

    def populated_columns(self) -> list[str]:
        """Columns holding at least one non-empty value."""
        return [c for c in self.columns if any(r.get(c) for r in self.rows)]

    @property
    def n_rows(self) -> int:
        return len(self.rows)


@dataclass
class Manifest:
    """A whole manifest workbook."""

    label: str
    provenance: Provenance
    sheets: dict[str, Sheet]

    def __getitem__(self, sheet: str) -> Sheet:
        return self.sheets[sheet]

    def get(self, sheet: str) -> Sheet | None:
        return self.sheets.get(sheet)

    @property
    def study_title(self) -> str:
        """The paper title, used to confirm two manifests describe one study."""
        study = self.sheets.get("study")
        if study and study.rows:
            return study.rows[0].get("project_name", "")
        return ""


def _clean(value: object) -> str:
    """Render a cell as a comparable string.

    openpyxl returns numbers as int/float regardless of how they were typed, which
    turns the instrument model 6000 into '6000.0' and an offset 0 into '0.0'.
    Integral floats are rendered without the misleading decimal tail.
    """
    if value is None:
        return ""
    if isinstance(value, bool):
        return "true" if value else "false"
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, (_dt.datetime, _dt.date)):
        return value.isoformat()
    return str(value).strip()


def read_provenance(path: Path) -> Provenance:
    """Hash a workbook and read its internal authoring metadata."""
    data = path.read_bytes()
    stat = path.stat()
    core: dict[str, str] = {}
    try:
        with zipfile.ZipFile(path) as archive:
            xml = archive.read(_CORE_XML).decode("utf-8", "replace")
        for key, tag in _CORE_FIELDS.items():
            match = re.search(rf"<{tag}[^>]*>(.*?)</{tag}>", xml)
            if match:
                core[key] = match.group(1)
    except (KeyError, zipfile.BadZipFile):
        pass
    return Provenance(
        path=str(path),
        sha256=hashlib.sha256(data).hexdigest(),
        size_bytes=len(data),
        mtime=_dt.datetime.fromtimestamp(stat.st_mtime).isoformat(timespec="seconds"),
        **core,
    )


def load_manifest(path: str | Path, label: str = "") -> Manifest:
    """Read a manifest workbook into a `Manifest`.

    Headers are normalised, wholly blank rows dropped, and cells rendered as
    strings. Sheets absent from the workbook are simply absent from the result;
    a missing sheet is a finding for the comparison, not an error here.
    """
    path = Path(path)
    workbook = load_workbook(path, data_only=True, read_only=True)
    sheets: dict[str, Sheet] = {}
    try:
        for worksheet in workbook.worksheets:
            rows = worksheet.iter_rows(values_only=True)
            try:
                header = next(rows)
            except StopIteration:
                continue
            columns: list[str] = []
            seen: set[str] = set()
            for raw in header:
                name = normalise_column(raw)
                if not name or name in seen:
                    # Blank or duplicated headers carry no identity; skip so a
                    # later column cannot silently overwrite an earlier one.
                    columns.append("")
                    continue
                seen.add(name)
                columns.append(name)
            records: list[dict[str, str]] = []
            for raw_row in rows:
                values = [_clean(v) for v in raw_row]
                if not any(values):
                    continue
                record = {
                    col: values[i]
                    for i, col in enumerate(columns)
                    if col and i < len(values) and values[i]
                }
                if record:
                    records.append(record)
            sheets[worksheet.title] = Sheet(
                name=worksheet.title,
                columns=[c for c in columns if c],
                rows=records,
            )
    finally:
        workbook.close()
    return Manifest(
        label=label or path.stem,
        provenance=read_provenance(path),
        sheets=sheets,
    )


def missing_sheets(manifest: Manifest) -> list[str]:
    """Compared sheets absent from this workbook.

    Excluded sheets are not reported: a workbook with no `file` sheet is not
    missing anything the comparison looks at.
    """
    return [s for s in SCORED_SHEETS if s not in manifest.sheets]
